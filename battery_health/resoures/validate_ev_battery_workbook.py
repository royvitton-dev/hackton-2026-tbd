from __future__ import annotations
import openpyxl, datetime, collections, json, math
from pathlib import Path

PATH = Path('/mnt/data/ev_battery_health_mock_data_10000_v2.xlsx')
wb=openpyxl.load_workbook(PATH,data_only=False)

def read_sheet(name):
    ws=wb[name]
    headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
    rows=[]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(headers,row)))
    return rows

vehicles=read_sheet('01_Vehicle_Master')
users=read_sheet('03_User_Profile_Mock')
sessions=read_sheet('04_Charge_Sessions_Raw')
veh={v['vehicleId']:v for v in vehicles}
user={u['userId']:u for u in users}

def usable_kwh(v):
    bg=v.get('batteryGrossKwh'); uf=v.get('usableFactor'); pv=v.get('packVoltage'); ah=v.get('batteryAh')
    if bg not in (None,'') and uf not in (None,''):
        return float(bg)*float(uf)
    if pv not in (None,'') and ah not in (None,'') and uf not in (None,''):
        return float(pv)*float(ah)/1000*float(uf)
    return None

errors=[]; warnings=[]; metrics={}
metrics['vehicle_count']=len(vehicles)
metrics['user_count']=len(users)
metrics['session_count']=len(sessions)

if len(vehicles) < 20: errors.append(('vehicle_count_lt_20', len(vehicles)))
if len(users) != 1250: errors.append(('user_count_not_1250', len(users)))
if len(sessions) != 10000: errors.append(('session_count_not_10000', len(sessions)))

# Duplicates
for name, rows, key in [('vehicle',vehicles,'vehicleId'),('user',users,'userId'),('session',sessions,'sessionId')]:
    counts=collections.Counter(r[key] for r in rows)
    dups=[k for k,c in counts.items() if c>1]
    metrics[f'duplicate_{name}_ids']=len(dups)
    if dups: errors.append((f'duplicate_{name}_ids', len(dups), dups[:10]))

# Mandatory vehicle fields for MVP
mandatory_vehicle=['vehicleId','manufacturer','modelName','modelYear','trimName','batteryGrossKwh','usableFactor','batteryChemistry','packVoltageClass','maxAcChargeKw','maxDcChargeKw','certifiedRangeKm','efficiencyKmPerKwh','sourceUrl','dataConfidence']
blank_vehicle=[]
for v in vehicles:
    for f in mandatory_vehicle:
        if v.get(f) in (None,''):
            blank_vehicle.append((v['vehicleId'],f))
metrics['blank_vehicle_mandatory']=len(blank_vehicle)
if blank_vehicle: errors.append(('blank_vehicle_mandatory', len(blank_vehicle), blank_vehicle[:10]))

# Session validity and distributions
by_user=collections.defaultdict(list)
unknown_user=[]; unknown_vehicle=[]; mismatch=[]; time_order=[]; overlap_end=[]; overlap_unplug=[]
charged_invalid=[]; c_rate_gt_4=[]; soc_bounds=[]; charger_counts=collections.Counter()
for s in sessions:
    sid=s['sessionId']; uid=s['userId']; vid=s['vehicleId']
    charger_counts[s['chargerType']]+=1
    if uid not in user:
        unknown_user.append(sid)
    else:
        if user[uid]['vehicleId'] != vid:
            mismatch.append(sid)
    if vid not in veh:
        unknown_vehicle.append(sid)
    st=s['startedAt']; en=s['endedAt']; un=s['unpluggedAt']
    if not (isinstance(st, datetime.datetime) and isinstance(en, datetime.datetime) and isinstance(un, datetime.datetime) and st < en <= un):
        time_order.append(sid)
    if s['chargedKwh'] is None or float(s['chargedKwh']) <= 0:
        charged_invalid.append(sid)
    if vid in veh and isinstance(st,datetime.datetime) and isinstance(en,datetime.datetime) and en>st:
        cap=usable_kwh(veh[vid])
        if cap:
            avg=float(s['chargedKwh'])/((en-st).total_seconds()/3600)
            cr=avg/cap
            if cr>4:
                c_rate_gt_4.append((sid,round(cr,2)))
    for sf in ['mockTruthStartSocPct','mockTruthEndSocPct']:
        val=s.get(sf)
        if val is not None and not (0 <= float(val) <= 100):
            soc_bounds.append((sid,sf,val))
    if uid in user:
        by_user[uid].append(s)

for uid,lst in by_user.items():
    lst=sorted(lst,key=lambda x:x['startedAt'])
    for a,b in zip(lst,lst[1:]):
        if b['startedAt'] < a['endedAt']:
            overlap_end.append((uid,a['sessionId'],b['sessionId']))
        if b['startedAt'] < a['unpluggedAt']:
            overlap_unplug.append((uid,a['sessionId'],b['sessionId']))

checks=[('unknown_user',unknown_user),('unknown_vehicle',unknown_vehicle),('user_vehicle_mismatch',mismatch),('invalid_time_order',time_order),('overlap_before_prev_end',overlap_end),('overlap_before_prev_unplug',overlap_unplug),('invalid_charged',charged_invalid),('soc_bounds',soc_bounds)]
for n,lst in checks:
    metrics[n]=len(lst)
    if lst: errors.append((n,len(lst),lst[:10]))
metrics['c_rate_gt_4_warning']=len(c_rate_gt_4)
if c_rate_gt_4: warnings.append(('c_rate_gt_4',len(c_rate_gt_4),c_rate_gt_4[:10]))
metrics['charger_counts']=dict(charger_counts)

# Eligibility independent calculation
eligible=0; insufficient=0; session_counts=[]; periods=[]; efcs=[]
for uid,u in user.items():
    lst=by_user.get(uid,[])
    cnt=len(lst); session_counts.append(cnt)
    if cnt:
        st=min(x['startedAt'] for x in lst); en=max(x['endedAt'] for x in lst)
        days=(en-st).total_seconds()/86400
        total=sum(float(x['chargedKwh']) for x in lst)
        cap=usable_kwh(veh[u['vehicleId']])
        efc=total/cap if cap else 0
    else:
        days=0; efc=0
    periods.append(days); efcs.append(efc)
    if cnt>=5 and days>=7 and efc>=0.3:
        eligible += 1
    else:
        insufficient += 1
metrics['eligible_users_calc']=eligible
metrics['insufficient_users_calc']=insufficient
metrics['sessions_per_user_min']=min(session_counts)
metrics['sessions_per_user_max']=max(session_counts)
metrics['sessions_per_user_avg']=round(sum(session_counts)/len(session_counts),3)
metrics['period_days_avg']=round(sum(periods)/len(periods),3)
metrics['efc_avg']=round(sum(efcs)/len(efcs),3)

# Formula sanity: no dynamic array functions, no unsupported obvious refs, formulas rows count.
formula_count=0; dynamic=[]; ref_errors=[]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v=cell.value
            if isinstance(v,str) and v.startswith('='):
                formula_count += 1
                upper=v.upper()
                for fn in ['FILTER(','XLOOKUP(','SORT(','SEQUENCE(','UNIQUE(','LET(','VSTACK(','HSTACK(']:
                    if fn in upper:
                        dynamic.append((ws.title,cell.coordinate,fn))
                if '#REF!' in upper:
                    ref_errors.append((ws.title,cell.coordinate,v))
metrics['formula_count']=formula_count
metrics['dynamic_formula_count']=len(dynamic)
metrics['ref_formula_count']=len(ref_errors)
if dynamic: errors.append(('dynamic_formulas',len(dynamic),dynamic[:5]))
if ref_errors: errors.append(('ref_errors_in_formulas',len(ref_errors),ref_errors[:5]))

# Output machine-readable report
report={'path':str(PATH),'errors':errors,'warnings':warnings,'metrics':metrics}
out=Path('/mnt/data/ev_battery_health_mock_data_10000_v2_validation.json')
out.write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2,default=str))
if errors:
    raise SystemExit(1)
